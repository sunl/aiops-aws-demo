"""Build the Word (and PDF) versions of the Markdown documents in this folder.

Supported Markdown subset: '#'..'####' headings, paragraphs, '- ' bullets and
'1. ' numbered items (nested by indentation), ``` code blocks (may be indented
under a list item), pipe tables, '> ' notes, '---' (ignored), **bold**,
`code`, [text](url), <br>, and ![caption](diagrams/x.png) on a line of its own.

Usage (from this folder):
    ../.venv-docs/bin/python build_docs.py              # every NN-*.md here -> .docx
    ../.venv-docs/bin/python build_docs.py --pdf        # also convert to PDF with LibreOffice
Diagrams: edit diagrams/*.svg, then run diagrams/export_png.sh before building.
"""
import glob
import os
import re
import subprocess
import sys

from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt

BODY_FONT = "宋体"
HEAD_FONT = "黑体"
CODE_FONT = "Courier New"   # Latin monospace; LibreOffice substitutes Liberation Mono
CODE_EA_FONT = "新宋体"      # Chinese characters inside code
PAGE_WIDTH_CM = 17.2  # A4 portrait, 1.9 cm side margins

INLINE = re.compile(r"(\*\*[^*]+\*\*|\[[^\]]+\]\([^)]+\)|`[^`]+`|<br>)")
IMAGE = re.compile(r"^!\[([^\]]*)\]\(([^)]+)\)\s*$")
LIST_ITEM = re.compile(r"^( *)(- |\d+\. )(.*)$")


def set_font(run, name, size=None, bold=None):
    run.font.name = name
    east = CODE_EA_FONT if name == CODE_FONT else name
    run._element.get_or_add_rPr().get_or_add_rFonts().set(qn("w:eastAsia"), east)
    if size is not None:
        run.font.size = Pt(size)
    if bold is not None:
        run.bold = bold


def add_hyperlink(paragraph, text, url, size):
    r_id = paragraph.part.relate_to(
        url, "http://schemas.openxmlformats.org/officeDocument/2006/relationships/hyperlink", is_external=True)
    link = OxmlElement("w:hyperlink")
    link.set(qn("r:id"), r_id)
    run = OxmlElement("w:r")
    rpr = OxmlElement("w:rPr")
    fonts = OxmlElement("w:rFonts")
    for attr in ("w:ascii", "w:hAnsi", "w:eastAsia"):
        fonts.set(qn(attr), BODY_FONT)
    rpr.append(fonts)
    color = OxmlElement("w:color")  # CT_RPr order: rFonts, color, sz, u
    color.set(qn("w:val"), "0563C1")
    rpr.append(color)
    sz = OxmlElement("w:sz")
    sz.set(qn("w:val"), str(int(size * 2)))
    rpr.append(sz)
    underline = OxmlElement("w:u")
    underline.set(qn("w:val"), "single")
    rpr.append(underline)
    run.append(rpr)
    t = OxmlElement("w:t")
    t.text = text
    t.set(qn("xml:space"), "preserve")
    run.append(t)
    link.append(run)
    paragraph._p.append(link)


def add_inline(paragraph, text, size, bold=False):
    for piece in INLINE.split(text):
        if not piece:
            continue
        if piece == "<br>":
            paragraph.add_run().add_break()
        elif piece.startswith("**") and piece.endswith("**") and len(piece) > 4:
            set_font(paragraph.add_run(piece[2:-2]), BODY_FONT, size, True)
        elif piece.startswith("`") and piece.endswith("`") and len(piece) > 2:
            set_font(paragraph.add_run(piece[1:-1]), CODE_FONT, size - 0.5, bold)
        elif piece.startswith("[") and re.match(r"\[([^\]]+)\]\(([^)]+)\)$", piece):
            m = re.match(r"\[([^\]]+)\]\(([^)]+)\)", piece)
            add_hyperlink(paragraph, m.group(1), m.group(2), size)
        else:
            set_font(paragraph.add_run(piece), BODY_FONT, size, bold)


def shade_paragraph(p, fill):
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:color"), "auto")
    shd.set(qn("w:fill"), fill)
    p._p.get_or_add_pPr().append(shd)


def shade_cell(cell, fill):
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:color"), "auto")
    shd.set(qn("w:fill"), fill)
    cell._tc.get_or_add_tcPr().append(shd)


def repeat_header(row):
    el = OxmlElement("w:tblHeader")
    el.set(qn("w:val"), "true")
    row._tr.get_or_add_trPr().append(el)


def split_row(line):
    return [c.strip() for c in line.strip().strip("|").split("|")]


def plain_len(text):
    text = re.sub(r"\[([^\]]+)\]\([^)]+\)", r"\1", text).replace("**", "").replace("`", "")
    # CJK characters are about twice as wide as Latin ones
    return max(sum(2 if ord(ch) > 0x2E80 else 1 for ch in part) for part in text.split("<br>"))


UNIT_CM = 0.165     # one Latin character at 9 pt (a CJK character counts as two units)
PADDING_CM = 0.35   # cell margins
SHORT_CM = 3.6      # columns whose text fits in this width keep their natural width


def auto_widths(rows):
    """Short columns get their natural width; long columns share the rest."""
    natural = []
    for c in range(len(rows[0])):
        longest = max(plain_len(r[c]) for r in rows if c < len(r))
        natural.append(max(longest, 4) * UNIT_CM + PADDING_CM)
    if sum(natural) <= PAGE_WIDTH_CM:
        scale = PAGE_WIDTH_CM / sum(natural)
        return [w * scale for w in natural]
    short = {c for c, w in enumerate(natural) if w <= SHORT_CM}
    rest = PAGE_WIDTH_CM - sum(natural[c] for c in short)
    long_cols = [c for c in range(len(natural)) if c not in short]
    if not long_cols or rest < 3.0 * len(long_cols):
        total = sum(min(w, 12) for w in natural)
        return [PAGE_WIDTH_CM * min(w, 12) / total for w in natural]
    weights = {c: min(natural[c], 12) for c in long_cols}  # dampen very long columns
    total = sum(weights.values())
    return [natural[c] if c in short else rest * weights[c] / total for c in range(len(natural))]


def keep_row_together(row):
    el = OxmlElement("w:cantSplit")
    el.set(qn("w:val"), "true")
    row._tr.get_or_add_trPr().append(el)


def add_table(doc, rows):
    header, body = rows[0], rows[1:]
    widths = auto_widths(rows)
    table = doc.add_table(rows=1 + len(body), cols=len(header))
    table.style = "Table Grid"
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    for c, w in enumerate(widths):  # LibreOffice lays tables out from the grid, Word from the cells
        table.columns[c].width = Cm(w)
    for r, values in enumerate([header] + body):
        row = table.rows[r]
        keep_row_together(row)
        for c in range(len(header)):
            cell = row.cells[c]
            cell.width = Cm(widths[c])
            p = cell.paragraphs[0]
            p.paragraph_format.space_after = Pt(0)
            add_inline(p, values[c] if c < len(values) else "", 9, bold=(r == 0))
            if r == 0:
                shade_cell(cell, "D9E2F3")
                p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    repeat_header(table.rows[0])
    doc.add_paragraph().paragraph_format.space_after = Pt(2)


def add_figure(doc, caption, path):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_after = Pt(2)
    p.paragraph_format.keep_with_next = True
    p.add_run().add_picture(path, width=Cm(PAGE_WIDTH_CM))
    if caption:
        c = doc.add_paragraph()
        c.alignment = WD_ALIGN_PARAGRAPH.CENTER
        c.paragraph_format.space_after = Pt(8)
        set_font(c.add_run(caption), BODY_FONT, 9)


def add_page_number(section):
    p = section.footer.paragraphs[0]
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    for tag, text in (("begin", None), (None, "PAGE"), ("end", None)):
        run = p.add_run()
        if tag:
            fld = OxmlElement("w:fldChar")
            fld.set(qn("w:fldCharType"), tag)
            run._r.append(fld)
        else:
            instr = OxmlElement("w:instrText")
            instr.set(qn("xml:space"), "preserve")
            instr.text = text
            run._r.append(instr)


def add_heading(doc, text, level):
    sizes = {1: 18, 2: 14, 3: 12, 4: 11}
    p = doc.add_paragraph()
    if level == 1:
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.paragraph_format.space_after = Pt(12)
    else:
        p.paragraph_format.space_before = Pt({2: 14, 3: 10, 4: 8}[level])
        p.paragraph_format.space_after = Pt(4 if level > 2 else 6)
        p.paragraph_format.keep_with_next = True
    set_font(p.add_run(text), HEAD_FONT, sizes[level], True)


def list_marker(marker, depth):
    if marker.strip().endswith("."):
        return marker.strip() + " "
    return ("• ", "◦ ", "▪ ")[min(depth, 2)]


def build(md_path, out_path):
    doc = Document()
    sec = doc.sections[0]
    sec.page_width, sec.page_height = Cm(21.0), Cm(29.7)
    sec.left_margin = sec.right_margin = Cm(1.9)
    sec.top_margin = sec.bottom_margin = Cm(2.0)
    normal = doc.styles["Normal"]
    normal.font.name = BODY_FONT
    normal.font.size = Pt(10.5)
    normal.element.rPr.rFonts.set(qn("w:eastAsia"), BODY_FONT)
    add_page_number(sec)
    base = os.path.dirname(os.path.abspath(md_path))

    lines = open(md_path, encoding="utf-8").read().splitlines()
    indents = []  # indentation of the open list levels
    i = 0
    while i < len(lines):
        line = lines[i]
        stripped = line.lstrip(" ")
        indent = len(line) - len(stripped)

        if stripped.startswith("```"):
            depth = sum(1 for x in indents if x < indent)
            i += 1
            while i < len(lines) and not lines[i].lstrip(" ").startswith("```"):
                text = lines[i][indent:] if lines[i][:indent].strip() == "" else lines[i].lstrip()
                p = doc.add_paragraph()
                shade_paragraph(p, "F2F2F2")  # shading before spacing (pPr order)
                p.paragraph_format.space_after = Pt(0)
                p.paragraph_format.left_indent = Cm(0.6 * depth)
                set_font(p.add_run(text or " "), CODE_FONT, 9)
                i += 1
            i += 1
            doc.add_paragraph().paragraph_format.space_after = Pt(2)
            continue

        image = IMAGE.match(stripped)
        if image:
            add_figure(doc, image.group(1), os.path.join(base, image.group(2)))
            i += 1
            continue

        if stripped.startswith("|"):
            block = []
            while i < len(lines) and lines[i].lstrip(" ").startswith("|"):
                block.append(lines[i].strip())
                i += 1
            rows = [split_row(l) for l in block if not re.match(r"^\|[\s\-|:]+\|$", l)]
            add_table(doc, rows)
            continue

        m = re.match(r"^(#{1,4}) (.*)$", line)
        if m:
            indents = []
            add_heading(doc, m.group(2), len(m.group(1)))
            i += 1
            continue

        if stripped.startswith("> "):
            p = doc.add_paragraph()
            shade_paragraph(p, "EEF4FF")
            p.paragraph_format.space_after = Pt(4)
            p.paragraph_format.left_indent = Cm(0.3)
            add_inline(p, stripped[2:], 10)
            i += 1
            continue

        item = LIST_ITEM.match(line)
        if item:
            while indents and indents[-1] > indent:
                indents.pop()
            if not indents or indents[-1] < indent:
                indents.append(indent)
            depth = len(indents) - 1
            p = doc.add_paragraph()
            p.paragraph_format.space_after = Pt(3)
            p.paragraph_format.left_indent = Cm(0.7 + 0.7 * depth)
            p.paragraph_format.first_line_indent = Cm(-0.5)
            add_inline(p, list_marker(item.group(2), depth) + item.group(3), 10.5)
            i += 1
            continue

        if stripped == "---" or not stripped:
            if not stripped and not (i + 1 < len(lines) and lines[i + 1].startswith(" ")):
                indents = []
            i += 1
            continue

        # plain paragraph (continuation text under a list item keeps its indent)
        depth = sum(1 for x in indents if x < indent) if indent else 0
        if not indent:
            indents = []
        p = doc.add_paragraph()
        p.paragraph_format.space_after = Pt(4)
        p.paragraph_format.left_indent = Cm(0.7 * depth)
        add_inline(p, stripped, 10.5)
        i += 1

    doc.save(out_path)


def to_pdf(docx_path):
    out_dir = os.path.dirname(os.path.abspath(docx_path))
    subprocess.run(["soffice", "--headless", "--convert-to", "pdf", "--outdir", out_dir, docx_path],
                   check=True, stdout=subprocess.DEVNULL, stderr=subprocess.DEVNULL, timeout=300)
    return os.path.splitext(docx_path)[0] + ".pdf"


if __name__ == "__main__":
    os.chdir(os.path.dirname(os.path.abspath(__file__)))
    want_pdf = "--pdf" in sys.argv
    sources = [a for a in sys.argv[1:] if a.endswith(".md")] or sorted(glob.glob("[0-9][0-9]-*.md"))
    for md in sources:
        out = os.path.splitext(md)[0] + ".docx"
        build(md, out)
        print("built", out)
        if want_pdf:
            print("built", to_pdf(out))
